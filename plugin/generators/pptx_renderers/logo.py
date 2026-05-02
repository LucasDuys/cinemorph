"""
Logo shape renderer for PPTX.

Renders a text box with brand mark or logo text in large, bold font.
Primarily for branding and wordmark display.

Spec: morph-deck-outputs R004.AC1 (shape renderers)
"""

from pptx.util import Pt, Emu
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor

from ._base import apply_morph_name, apply_opacity, RendererCtx
try:
    from lib.layout import parse_layout
except ImportError:
    from ..lib.layout import parse_layout


def render(stage_idx, element_id, layout, props, slide, ctx) -> None:
    """
    Render a logo text box.

    Args:
        stage_idx: int, 0-indexed stage number
        element_id: str, unique element ID (e.g. "logo_1")
        layout: dict, position + size in percent (left, top, width, height)
        props: dict, element properties:
            - text: str, logo or brand mark text
            - color: str (hex), optional text color; defaults to theme foreground
            - font_size: int, optional font size in points; defaults to 32
        slide: pptx.shapes.slide.Slide, target slide
        ctx: RendererCtx, renderer context (theme, fonts, etc.)
    """
    # Parse layout to EMU coordinates
    parsed = parse_layout(layout)
    left = parsed['left']
    top = parsed['top']
    width = parsed['width']
    height = parsed['height']

    # Add text box
    text_box = slide.shapes.add_textbox(left, top, width, height)

    # Apply morph name for tracking
    apply_morph_name(text_box, element_id)

    # Configure text frame
    tf = text_box.text_frame
    tf.word_wrap = True
    tf.margin_left = Emu(60000)
    tf.margin_right = Emu(60000)
    tf.margin_top = Emu(60000)
    tf.margin_bottom = Emu(60000)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE

    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = ""

    # Get logo text
    logo_text = props.get('text', '')

    # Add run with styling
    run = p.add_run()
    run.text = logo_text

    # Set font size (large and bold for brand display)
    font_size = props.get('font_size', 32)
    try:
        font_size = int(font_size)
    except (ValueError, TypeError):
        font_size = 32
    run.font.size = Pt(font_size)
    run.font.bold = True

    # Determine text color
    text_color = None
    if 'color' in props and props['color']:
        hex_color = props['color'].lstrip('#')
        try:
            text_color = RGBColor(
                int(hex_color[0:2], 16),
                int(hex_color[2:4], 16),
                int(hex_color[4:6], 16)
            )
        except (ValueError, IndexError):
            text_color = None

    if text_color is None:
        # Use theme foreground
        foreground = ctx.theme.get('foreground', RGBColor(0xFA, 0xFA, 0xFA))
        if isinstance(foreground, str):
            hex_color = foreground.lstrip('#')
            text_color = RGBColor(
                int(hex_color[0:2], 16),
                int(hex_color[2:4], 16),
                int(hex_color[4:6], 16)
            )
        else:
            text_color = foreground

    run.font.color.rgb = text_color
    run.font.name = ctx.fonts.get('display', 'Inter')

    # Apply opacity if present
    if 'opacity' in props:
        try:
            apply_opacity(text_box, float(props['opacity']))
        except (ValueError, TypeError):
            pass  # Ignore invalid opacity


RENDERER_REGISTRY = {}
RENDERER_REGISTRY['logo'] = render
