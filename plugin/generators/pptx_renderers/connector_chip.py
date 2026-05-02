"""
Connector chip shape renderer for PPTX.

Renders connector tiles with brand-fill rounded rectangles and text labels.
Supports multiple layout variants: orbit (icon-only square), cluster (icon-only square),
pipeline (wider pill with text), footer (small square).

Spec: morph-deck-outputs R004.AC1 (shape renderers)
"""

from pptx.util import Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor

from ._base import apply_morph_name, apply_opacity, RendererCtx
try:
    from lib.layout import parse_layout
except ImportError:
    from ..lib.layout import parse_layout


def render(stage_idx, element_id, layout, props, slide, ctx) -> None:
    """
    Render a connector chip shape.

    Args:
        stage_idx: int, 0-indexed stage number
        element_id: str, unique element ID (e.g. "connector_chip_drive")
        layout: dict, position + size in percent (left, top, width, height)
        props: dict, element properties:
            - label: str, text label (e.g. "Drive")
            - brand_color: str (hex), optional explicit color; falls back to theme palette
            - id: str, connector ID to look up in theme connectorPalette
            - layout.shape: str, layout variant ("orbit", "cluster", "pipeline", "footer")
        slide: pptx.shapes.slide.Slide, target slide
        ctx: RendererCtx, renderer context (theme, fonts, etc.)
    """
    # Parse layout to EMU coordinates
    parsed = parse_layout(layout)
    left = parsed['left']
    top = parsed['top']
    width = parsed['width']
    height = parsed['height']

    # Determine fill color: explicit brand_color -> theme connectorPalette -> fallback
    fill_color = None
    if 'brand_color' in props and props['brand_color']:
        # Parse hex color
        hex_color = props['brand_color']
        if hex_color.startswith('#'):
            hex_color = hex_color[1:]
        try:
            fill_color = RGBColor(
                int(hex_color[0:2], 16),
                int(hex_color[2:4], 16),
                int(hex_color[4:6], 16)
            )
        except (ValueError, IndexError):
            fill_color = None

    if fill_color is None and 'id' in props:
        # Try theme connector palette
        connector_id = props['id']
        if 'connectorPalette' in ctx.theme and connector_id in ctx.theme['connectorPalette']:
            palette_entry = ctx.theme['connectorPalette'][connector_id]
            if isinstance(palette_entry, str):
                # Hex string
                hex_color = palette_entry.lstrip('#')
                try:
                    fill_color = RGBColor(
                        int(hex_color[0:2], 16),
                        int(hex_color[2:4], 16),
                        int(hex_color[4:6], 16)
                    )
                except (ValueError, IndexError):
                    pass

    if fill_color is None:
        # Fallback: use theme primary or default blue
        primary = ctx.theme.get('primary', '#4285F4')
        hex_color = primary.lstrip('#')
        try:
            fill_color = RGBColor(
                int(hex_color[0:2], 16),
                int(hex_color[2:4], 16),
                int(hex_color[4:6], 16)
            )
        except (ValueError, IndexError):
            fill_color = RGBColor(0x42, 0x85, 0xF4)  # fallback blue

    # Add rounded rectangle shape
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        left, top, width, height
    )

    # Apply morph name for tracking
    apply_morph_name(shape, element_id)

    # Set fill color
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color

    # Set border: thin, subtle
    border_color = ctx.theme.get('border', RGBColor(0x27, 0x27, 0x2A))
    if isinstance(border_color, str):
        border_hex = border_color.lstrip('#')
        border_color = RGBColor(
            int(border_hex[0:2], 16),
            int(border_hex[2:4], 16),
            int(border_hex[4:6], 16)
        )
    shape.line.color.rgb = border_color
    shape.line.width = Pt(0.5)

    # Set corner radius (rounded rectangle adjustment)
    shape.adjustments[0] = 0.15

    # Remove shadow
    shape.shadow.inherit = False

    # Get layout shape variant to determine label behavior
    layout_shape = layout.get('shape', 'orbit')

    # Add text label (only for pipeline, footer; icon-only for orbit/cluster)
    if layout_shape in ('pipeline', 'footer'):
        tf = shape.text_frame
        tf.word_wrap = False
        tf.margin_left = tf.margin_right = 0
        tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE

        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.text = ""

        label = props.get('label', '')
        run = p.add_run()
        run.text = label
        run.font.size = Pt(10)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0xFA, 0xFA, 0xFA)  # foreground
        run.font.name = ctx.fonts.get('body', 'Inter')

    # Apply opacity if present
    if 'opacity' in props:
        try:
            apply_opacity(shape, float(props['opacity']))
        except (ValueError, TypeError):
            pass  # Ignore invalid opacity


RENDERER_REGISTRY = {}
RENDERER_REGISTRY['connector_chip'] = render
