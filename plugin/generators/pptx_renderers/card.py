"""
Card shape renderer for PPTX.

Renders a rounded rectangle with subtle border, serving as a container for
content. Supports text fallback via children_text property when no nested
elements are rendered (group renderer responsibility).

Spec: morph-deck-outputs R004.AC1 (shape renderers)
"""

from pptx.util import Pt, Emu
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor

from ._base import apply_morph_name, apply_opacity, RendererCtx
try:
    from lib.layout import parse_layout
except ImportError:
    from ..lib.layout import parse_layout


def render(stage_idx, element_id, layout, props, slide, ctx) -> None:
    """
    Render a card shape.

    Args:
        stage_idx: int, 0-indexed stage number
        element_id: str, unique element ID (e.g. "card_1")
        layout: dict, position + size in percent (left, top, width, height)
        props: dict, element properties:
            - children_text: str, optional fallback text if no nested elements
            - bg_color: str (hex), optional background; defaults to theme surface
        slide: pptx.shapes.slide.Slide, target slide
        ctx: RendererCtx, renderer context (theme, fonts, etc.)
    """
    # Parse layout to EMU coordinates
    parsed = parse_layout(layout)
    left = parsed['left']
    top = parsed['top']
    width = parsed['width']
    height = parsed['height']

    # Determine background color
    bg_color = None
    if 'bg_color' in props and props['bg_color']:
        hex_color = props['bg_color'].lstrip('#')
        try:
            bg_color = RGBColor(
                int(hex_color[0:2], 16),
                int(hex_color[2:4], 16),
                int(hex_color[4:6], 16)
            )
        except (ValueError, IndexError):
            bg_color = None

    if bg_color is None:
        # Use theme surface color
        surface = ctx.theme.get('surface', RGBColor(0x1D, 0x1D, 0x21))
        if isinstance(surface, str):
            hex_color = surface.lstrip('#')
            bg_color = RGBColor(
                int(hex_color[0:2], 16),
                int(hex_color[2:4], 16),
                int(hex_color[4:6], 16)
            )
        else:
            bg_color = surface

    # Add rounded rectangle shape
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        left, top, width, height
    )

    # Apply morph name for tracking
    apply_morph_name(shape, element_id)

    # Set background fill
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_color

    # Set subtle border
    border_color = ctx.theme.get('border', RGBColor(0x27, 0x27, 0x2A))
    if isinstance(border_color, str):
        border_hex = border_color.lstrip('#')
        border_color = RGBColor(
            int(border_hex[0:2], 16),
            int(border_hex[2:4], 16),
            int(border_hex[4:6], 16)
        )
    shape.line.color.rgb = border_color
    shape.line.width = Pt(0.75)

    # Set corner radius
    shape.adjustments[0] = 0.12

    # Remove shadow
    shape.shadow.inherit = False

    # Add fallback text if children_text is provided
    if 'children_text' in props and props['children_text']:
        tf = shape.text_frame
        tf.word_wrap = True
        tf.margin_left = Emu(120000)
        tf.margin_right = Emu(120000)
        tf.margin_top = Emu(60000)
        tf.margin_bottom = Emu(60000)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE

        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        p.text = ""

        run = p.add_run()
        run.text = props['children_text']
        run.font.size = Pt(12)
        run.font.color.rgb = ctx.theme.get('foreground', RGBColor(0xFA, 0xFA, 0xFA))
        run.font.name = ctx.fonts.get('body', 'Inter')

    # Apply opacity if present
    if 'opacity' in props:
        try:
            apply_opacity(shape, float(props['opacity']))
        except (ValueError, TypeError):
            pass  # Ignore invalid opacity


RENDERER_REGISTRY = {}
RENDERER_REGISTRY['card'] = render
