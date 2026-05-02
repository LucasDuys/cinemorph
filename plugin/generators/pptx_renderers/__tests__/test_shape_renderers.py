"""
Tests for PPTX shape primitive renderers.

Tests cover:
  - connector_chip: rounded rect with brand fill + label
  - card: rounded rect with subtle border
  - logo: textbox with large bold text
  - image: image load or placeholder fallback
  - footer_strip: thin rectangle bar at bottom

Spec: morph-deck-outputs R004.AC1 (shape renderers)
"""

import unittest
from pathlib import Path
from unittest.mock import Mock, MagicMock, patch
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor

from generators.pptx_renderers._base import RendererCtx, apply_morph_name, apply_opacity


class TestConnectorChipRenderer(unittest.TestCase):
    """Test connector_chip renderer."""

    def setUp(self):
        """Create a presentation and slide for testing."""
        self.prs = Presentation()
        self.slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])  # blank layout
        self.ctx = RendererCtx(
            theme={
                'primary': '#4285F4',
                'border': '#27272A',
                'connectorPalette': {
                    'drive': '#4285F4',
                    'slack': '#ECB22E',
                },
            },
            dwell_ms=1000,
            slide_idx=0,
            fonts={'body': 'Inter'},
            dist_dir=Path('/tmp'),
        )

    def test_connector_chip_basic(self):
        """Test connector_chip renders shape with brand fill."""
        from generators.pptx_renderers.connector_chip import render

        layout = {'left': '10%', 'top': '10%', 'width': '5%', 'height': '5%'}
        props = {'label': 'Drive', 'id': 'drive'}

        render(0, 'connector_chip_drive', layout, props, self.slide, self.ctx)

        self.assertGreater(len(self.slide.shapes), 0)
        shape = self.slide.shapes[-1]
        self.assertTrue(shape.name.startswith('!!'))

    def test_connector_chip_with_explicit_color(self):
        """Test connector_chip with explicit brand_color."""
        from generators.pptx_renderers.connector_chip import render

        layout = {'left': '20%', 'top': '20%', 'width': '5%', 'height': '5%'}
        props = {'label': 'Custom', 'brand_color': '#FF0000'}

        render(0, 'connector_chip_custom', layout, props, self.slide, self.ctx)

        shape = self.slide.shapes[-1]
        self.assertTrue(shape.name.startswith('!!'))
        self.assertTrue(shape.fill.type == 1)  # SOLID

    def test_connector_chip_pipeline_layout(self):
        """Test connector_chip with pipeline layout variant."""
        from generators.pptx_renderers.connector_chip import render

        layout = {
            'left': '30%', 'top': '30%', 'width': '10%', 'height': '3%',
            'shape': 'pipeline'
        }
        props = {'label': 'Pipeline', 'id': 'slack'}

        render(0, 'connector_chip_pipeline', layout, props, self.slide, self.ctx)

        shape = self.slide.shapes[-1]
        self.assertIsNotNone(shape.text_frame)


class TestCardRenderer(unittest.TestCase):
    """Test card renderer."""

    def setUp(self):
        """Create a presentation and slide for testing."""
        self.prs = Presentation()
        self.slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.ctx = RendererCtx(
            theme={
                'surface': '#1D1D21',
                'border': '#27272A',
            },
            dwell_ms=1000,
            slide_idx=0,
            fonts={'body': 'Inter'},
            dist_dir=Path('/tmp'),
        )

    def test_card_basic(self):
        """Test card renders rounded rectangle."""
        from generators.pptx_renderers.card import render

        layout = {'left': '10%', 'top': '40%', 'width': '20%', 'height': '15%'}
        props = {}

        render(0, 'card_1', layout, props, self.slide, self.ctx)

        shape = self.slide.shapes[-1]
        self.assertTrue(shape.name.startswith('!!'))
        self.assertTrue(shape.fill.type == 1)  # SOLID

    def test_card_with_fallback_text(self):
        """Test card with children_text fallback."""
        from generators.pptx_renderers.card import render

        layout = {'left': '50%', 'top': '40%', 'width': '20%', 'height': '15%'}
        props = {'children_text': 'Fallback content'}

        render(0, 'card_2', layout, props, self.slide, self.ctx)

        shape = self.slide.shapes[-1]
        self.assertIsNotNone(shape.text_frame)
        self.assertIn('Fallback', shape.text_frame.text)


class TestLogoRenderer(unittest.TestCase):
    """Test logo renderer."""

    def setUp(self):
        """Create a presentation and slide for testing."""
        self.prs = Presentation()
        self.slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.ctx = RendererCtx(
            theme={
                'foreground': '#FAFAFA',
            },
            dwell_ms=1000,
            slide_idx=0,
            fonts={'display': 'Inter'},
            dist_dir=Path('/tmp'),
        )

    def test_logo_basic(self):
        """Test logo renders textbox with bold text."""
        from generators.pptx_renderers.logo import render

        layout = {'left': '30%', 'top': '5%', 'width': '40%', 'height': '10%'}
        props = {'text': 'STACKLINK'}

        render(0, 'logo_1', layout, props, self.slide, self.ctx)

        shape = self.slide.shapes[-1]
        self.assertTrue(shape.name.startswith('!!'))
        self.assertIn('STACKLINK', shape.text_frame.text)

    def test_logo_with_custom_size(self):
        """Test logo with custom font_size."""
        from generators.pptx_renderers.logo import render

        layout = {'left': '30%', 'top': '5%', 'width': '40%', 'height': '10%'}
        props = {'text': 'Brand', 'font_size': 48}

        render(0, 'logo_2', layout, props, self.slide, self.ctx)

        shape = self.slide.shapes[-1]
        self.assertIsNotNone(shape.text_frame.paragraphs[0].runs)


class TestImageRenderer(unittest.TestCase):
    """Test image renderer."""

    def setUp(self):
        """Create a presentation and slide for testing."""
        self.prs = Presentation()
        self.slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.ctx = RendererCtx(
            theme={
                'muted_foreground': '#A1A1AA',
                'border': '#27272A',
            },
            dwell_ms=1000,
            slide_idx=0,
            fonts={'body': 'Inter'},
            dist_dir=Path('/tmp'),
        )

    def test_image_placeholder_fallback(self):
        """Test image renders placeholder when file not found."""
        from generators.pptx_renderers.image import render

        layout = {'left': '10%', 'top': '60%', 'width': '20%', 'height': '15%'}
        props = {'src': '/nonexistent/image.png', 'alt': 'Missing Image'}

        render(0, 'image_1', layout, props, self.slide, self.ctx)

        shape = self.slide.shapes[-1]
        self.assertTrue(shape.name.startswith('!!'))
        # Placeholder should have alt text
        self.assertIn('Missing Image', shape.text_frame.text)

    def test_image_no_src(self):
        """Test image with no src renders placeholder."""
        from generators.pptx_renderers.image import render

        layout = {'left': '40%', 'top': '60%', 'width': '20%', 'height': '15%'}
        props = {'alt': 'Default Alt'}

        render(0, 'image_2', layout, props, self.slide, self.ctx)

        shape = self.slide.shapes[-1]
        self.assertTrue(shape.name.startswith('!!'))


class TestFooterStripRenderer(unittest.TestCase):
    """Test footer_strip renderer."""

    def setUp(self):
        """Create a presentation and slide for testing."""
        self.prs = Presentation()
        self.slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.ctx = RendererCtx(
            theme={
                'surface': '#1D1D21',
                'muted_foreground': '#A1A1AA',
            },
            dwell_ms=1000,
            slide_idx=0,
            fonts={'body': 'Inter'},
            dist_dir=Path('/tmp'),
        )

    def test_footer_strip_basic(self):
        """Test footer_strip renders thin rectangle."""
        from generators.pptx_renderers.footer_strip import render

        layout = {'left': '0%', 'top': '92%', 'width': '100%', 'height': '8%'}
        props = {}

        render(0, 'footer_strip_1', layout, props, self.slide, self.ctx)

        shape = self.slide.shapes[-1]
        self.assertTrue(shape.name.startswith('!!'))
        self.assertTrue(shape.fill.type == 1)  # SOLID

    def test_footer_strip_with_text(self):
        """Test footer_strip with optional text."""
        from generators.pptx_renderers.footer_strip import render

        layout = {'left': '0%', 'top': '92%', 'width': '100%', 'height': '8%'}
        props = {'text': 'Slide 1 of 5'}

        render(0, 'footer_strip_2', layout, props, self.slide, self.ctx)

        shape = self.slide.shapes[-1]
        self.assertIn('Slide', shape.text_frame.text)


if __name__ == '__main__':
    unittest.main()
