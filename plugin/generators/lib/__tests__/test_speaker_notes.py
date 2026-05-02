"""Tests for PPTX speaker notes formatting.

Spec: spec-morph-deck-outputs.md R005 (PPTX speaker notes from talkTrack).

Tests cover:
  - format_notes with full talkTrack (script + cues + dwell)
  - format_notes with cues only
  - format_notes returns empty when no script
  - BACKUP prefix on backup stages
  - dwell prefix matches stage.dwellSeconds
  - apply_notes_to_slide writes to actual slide.notes_slide.notes_text_frame
  - apply_notes_to_all validates lengths match
"""

import sys
from pathlib import Path

import pytest
from pptx import Presentation

# Add lib dir to path
sys.path.insert(0, str(Path(__file__).parent.parent))

from speaker_notes import format_notes, apply_notes_to_slide, apply_notes_to_all


class TestFormatNotes:
    """Test format_notes function."""

    def test_format_notes_with_full_talktrack(self):
        """Full talkTrack with script, cues, and dwell."""
        stage = {
            "talkTrack": {
                "script": "This is the main talking point.",
                "dwellSeconds": 8,
                "cues": ["Pause here", "Then press next"],
            }
        }
        notes = format_notes(stage)

        assert "[~8s dwell]" in notes
        assert "This is the main talking point." in notes
        assert "Cues:" in notes
        assert "- Pause here" in notes
        assert "- Then press next" in notes

    def test_format_notes_with_cues_only(self):
        """TalkTrack with script and cues, no dwell."""
        stage = {
            "talkTrack": {
                "script": "Main point here.",
                "cues": ["Cue one", "Cue two"],
            }
        }
        notes = format_notes(stage)

        assert "Main point here." in notes
        assert "Cues:" in notes
        assert "- Cue one" in notes
        assert "- Cue two" in notes
        assert "dwell" not in notes

    def test_format_notes_script_only(self):
        """TalkTrack with only script, no cues or dwell."""
        stage = {
            "talkTrack": {
                "script": "Just the script.",
            }
        }
        notes = format_notes(stage)

        assert notes == "Just the script."
        assert "Cues:" not in notes

    def test_format_notes_empty_when_no_talktrack(self):
        """Returns empty string when stage has no talkTrack."""
        stage = {}
        notes = format_notes(stage)

        assert notes == ""

    def test_format_notes_empty_when_no_script(self):
        """Returns empty string when talkTrack exists but script is empty/missing."""
        stage = {"talkTrack": {"cues": ["orphan cue"]}}
        notes = format_notes(stage)

        assert notes == ""

    def test_format_notes_backup_prefix(self):
        """[BACKUP] prefix appears when isBackup=True."""
        stage = {
            "isBackup": True,
            "talkTrack": {
                "script": "Backup Q&A response.",
                "dwellSeconds": 5,
            },
        }
        notes = format_notes(stage)

        assert "[BACKUP]" in notes
        assert "[~5s dwell]" in notes
        assert "Backup Q&A response." in notes

    def test_format_notes_backup_without_dwell(self):
        """[BACKUP] prefix still appears when no dwell."""
        stage = {
            "isBackup": True,
            "talkTrack": {
                "script": "Backup response.",
            },
        }
        notes = format_notes(stage)

        assert "[BACKUP]" in notes
        assert "Backup response." in notes

    def test_format_notes_dwell_integer_cast(self):
        """dwellSeconds is cast to int (e.g., 7.5 -> 7)."""
        stage = {
            "talkTrack": {
                "script": "Test.",
                "dwellSeconds": 7.8,
            }
        }
        notes = format_notes(stage)

        assert "[~7s dwell]" in notes

    def test_format_notes_dwell_zero(self):
        """dwellSeconds=0 is valid and included."""
        stage = {
            "talkTrack": {
                "script": "Test.",
                "dwellSeconds": 0,
            }
        }
        notes = format_notes(stage)

        assert "[~0s dwell]" in notes

    def test_format_notes_empty_cues_list_ignored(self):
        """Empty cues list does not create 'Cues:' header."""
        stage = {
            "talkTrack": {
                "script": "Main point.",
                "cues": [],
            }
        }
        notes = format_notes(stage)

        assert "Cues:" not in notes
        assert notes == "Main point."

    def test_format_notes_multiline_script(self):
        """Multiline script is preserved."""
        stage = {
            "talkTrack": {
                "script": "Line one.\nLine two.\nLine three.",
            }
        }
        notes = format_notes(stage)

        assert "Line one.\nLine two.\nLine three." in notes

    def test_format_notes_script_with_leading_trailing_whitespace(self):
        """Script is trimmed of leading/trailing whitespace."""
        stage = {
            "talkTrack": {
                "script": "  \n  Main point.  \n  ",
            }
        }
        notes = format_notes(stage)

        assert notes == "Main point."

    def test_format_notes_multiple_cues(self):
        """Multiple cues are all listed."""
        stage = {
            "talkTrack": {
                "script": "Point.",
                "cues": ["First", "Second", "Third"],
            }
        }
        notes = format_notes(stage)

        assert "- First" in notes
        assert "- Second" in notes
        assert "- Third" in notes


class TestApplyNotesToSlide:
    """Test apply_notes_to_slide with real python-pptx Presentation."""

    def _make_presentation_with_slides(self, n: int):
        """Create a real Presentation with n blank slides."""
        prs = Presentation()
        blank_layout = prs.slide_layouts[6]  # blank layout
        slides = []
        for _ in range(n):
            slide = prs.slides.add_slide(blank_layout)
            slides.append(slide)
        return prs, slides

    def test_apply_notes_to_slide_writes_to_notes_frame(self):
        """Notes text is written to slide.notes_slide.notes_text_frame.text."""
        prs, slides = self._make_presentation_with_slides(1)
        slide = slides[0]

        stage = {
            "talkTrack": {
                "script": "Test speaker note.",
                "dwellSeconds": 6,
            }
        }

        apply_notes_to_slide(slide, stage)

        # Read back from the slide
        notes_text = slide.notes_slide.notes_text_frame.text
        assert "[~6s dwell]" in notes_text
        assert "Test speaker note." in notes_text

    def test_apply_notes_to_slide_with_cues(self):
        """Notes with cues are correctly written to slide."""
        prs, slides = self._make_presentation_with_slides(1)
        slide = slides[0]

        stage = {
            "talkTrack": {
                "script": "Main point.",
                "cues": ["First cue", "Second cue"],
            }
        }

        apply_notes_to_slide(slide, stage)

        notes_text = slide.notes_slide.notes_text_frame.text
        assert "Main point." in notes_text
        assert "Cues:" in notes_text
        assert "- First cue" in notes_text
        assert "- Second cue" in notes_text

    def test_apply_notes_to_slide_skips_empty_notes(self):
        """Empty notes do not write to the slide (no script)."""
        prs, slides = self._make_presentation_with_slides(1)
        slide = slides[0]

        # Stage with no talkTrack
        stage = {}

        apply_notes_to_slide(slide, stage)

        # notes_slide.notes_text_frame is created but empty
        notes_text = slide.notes_slide.notes_text_frame.text
        assert notes_text == ""

    def test_apply_notes_to_slide_with_backup_flag(self):
        """Backup stage notes include [BACKUP] prefix."""
        prs, slides = self._make_presentation_with_slides(1)
        slide = slides[0]

        stage = {
            "isBackup": True,
            "talkTrack": {
                "script": "Q&A backup.",
            }
        }

        apply_notes_to_slide(slide, stage)

        notes_text = slide.notes_slide.notes_text_frame.text
        assert "[BACKUP]" in notes_text


class TestApplyNotesToAll:
    """Test apply_notes_to_all with slide/stage pairs."""

    def _make_presentation_with_slides(self, n: int):
        """Create a real Presentation with n blank slides."""
        prs = Presentation()
        blank_layout = prs.slide_layouts[6]
        slides = []
        for _ in range(n):
            slide = prs.slides.add_slide(blank_layout)
            slides.append(slide)
        return prs, slides

    def test_apply_notes_to_all_paired_iteration(self):
        """All slides receive notes from their corresponding stages."""
        prs, slides = self._make_presentation_with_slides(3)

        stages = [
            {
                "talkTrack": {
                    "script": "First slide.",
                    "dwellSeconds": 5,
                }
            },
            {
                "talkTrack": {
                    "script": "Second slide.",
                }
            },
            {
                "talkTrack": {
                    "script": "Third slide.",
                    "cues": ["Ready"],
                }
            },
        ]

        apply_notes_to_all(slides, stages)

        # Verify each slide has correct notes
        assert "[~5s dwell]" in slides[0].notes_slide.notes_text_frame.text
        assert "First slide." in slides[0].notes_slide.notes_text_frame.text

        assert "Second slide." in slides[1].notes_slide.notes_text_frame.text

        assert "Third slide." in slides[2].notes_slide.notes_text_frame.text
        assert "- Ready" in slides[2].notes_slide.notes_text_frame.text

    def test_apply_notes_to_all_raises_on_length_mismatch(self):
        """Raises ValueError if slides and stages lengths don't match."""
        prs, slides = self._make_presentation_with_slides(2)
        stages = [{"talkTrack": {"script": "Only one."}}]

        with pytest.raises(ValueError, match="same length"):
            apply_notes_to_all(slides, stages)

    def test_apply_notes_to_all_empty_lists(self):
        """Empty slides and stages lists are valid (no-op)."""
        apply_notes_to_all([], [])
        # Should not raise

    def test_apply_notes_to_all_with_mixed_empty_stages(self):
        """Some stages with no talkTrack are skipped gracefully."""
        prs, slides = self._make_presentation_with_slides(2)
        stages = [
            {"talkTrack": {"script": "First."}},
            {},  # No talkTrack
        ]

        apply_notes_to_all(slides, stages)

        assert "First." in slides[0].notes_slide.notes_text_frame.text
        assert slides[1].notes_slide.notes_text_frame.text == ""
